"""Tests for Phase 5 Excel report generation."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from jira_analytics.reports.excel_builder import SHEET_NAMES, ExcelReportBuilder
from jira_analytics.reports.loader import load_report_inputs

REPO_ROOT = Path(__file__).resolve().parents[1]
CACHE_DIR = REPO_ROOT / "cache"


@pytest.fixture(scope="module")
def report_path(tmp_path_factory) -> Path:
    if not (CACHE_DIR / "metrics.json").exists():
        pytest.skip("cache/metrics.json not found — run pipeline first")
    inputs = load_report_inputs(CACHE_DIR)
    out = tmp_path_factory.mktemp("reports") / "test_report.xlsx"
    return ExcelReportBuilder(inputs).build(out)


class TestReportGeneration:
    def test_workbook_exists(self, report_path: Path) -> None:
        assert report_path.exists()
        assert report_path.stat().st_size > 0

    def test_workbook_opens(self, report_path: Path) -> None:
        wb = load_workbook(report_path)
        assert wb is not None
        wb.close()

    def test_all_sheets_exist(self, report_path: Path) -> None:
        wb = load_workbook(report_path)
        assert wb.sheetnames == SHEET_NAMES
        for name in SHEET_NAMES:
            ws = wb[name]
            assert ws.max_row >= 1
        wb.close()

    def test_charts_exist(self, report_path: Path) -> None:
        wb = load_workbook(report_path)
        chart_count = sum(len(ws._charts) for ws in wb.worksheets)
        assert chart_count >= 5, f"Expected at least 5 charts, found {chart_count}"
        wb.close()

    def test_totals_match_analytics(self, report_path: Path) -> None:
        inputs = load_report_inputs(CACHE_DIR)
        wb = load_workbook(report_path, data_only=True)
        throughput_ws = wb["Throughput"]

        total_in_sheet = None
        for row in throughput_ws.iter_rows(min_row=1, max_col=2, values_only=True):
            if row[0] == "Total Throughput":
                total_in_sheet = row[1]
                break

        assert total_in_sheet == inputs.analytics.summary.total_throughput

        exec_ws = wb["Executive Summary"]
        open_issues = None
        for row in exec_ws.iter_rows(min_row=1, max_col=3, values_only=True):
            if row[1] == "Open Issues":
                open_issues = row[2]
                break
        assert open_issues == inputs.analytics.summary.open_issues
        wb.close()

    def test_lead_time_columns(self, report_path: Path) -> None:
        wb = load_workbook(report_path)
        ws = wb["Lead Time"]
        headers = [cell.value for cell in ws[1]]
        expected = [
            "Issue Key",
            "Summary",
            "Lead Time (h)",
            "Cycle Time (h)",
            "Net Cycle Time (h)",
            "Buffer Time (h)",
            "Blocked Time (h)",
            "Time to First Progress (h)",
            "Waiting Time (h)",
            "Active Time (h)",
            "Reopened",
            "Terminal Status",
        ]
        assert headers == expected
        wb.close()

    def test_management_insights_not_empty(self, report_path: Path) -> None:
        wb = load_workbook(report_path)
        ws = wb["Management Insights"]
        assert ws.max_row >= 5
        wb.close()
