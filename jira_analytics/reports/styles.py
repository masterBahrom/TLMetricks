"""Shared openpyxl styles for the Excel report."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Brand palette
NAVY = "1F3864"
BLUE = "2F5496"
LIGHT_BLUE = "D6E4F0"
WHITE = "FFFFFF"
ALT_ROW = "F2F7FB"
GREEN = "C6EFCE"
YELLOW = "FFEB9C"
RED = "FFC7CE"
HEADER_FILL = PatternFill("solid", fgColor=BLUE)
TITLE_FILL = PatternFill("solid", fgColor=NAVY)
ALT_FILL = PatternFill("solid", fgColor=ALT_ROW)
KPI_GOOD = PatternFill("solid", fgColor=GREEN)
KPI_WARN = PatternFill("solid", fgColor=YELLOW)
KPI_BAD = PatternFill("solid", fgColor=RED)

THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TITLE_FONT = Font(name="Calibri", size=18, bold=True, color=WHITE)
SUBTITLE_FONT = Font(name="Calibri", size=11, color="595959")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
BODY_FONT = Font(name="Calibri", size=11)
KPI_LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="595959")
KPI_VALUE_FONT = Font(name="Calibri", size=20, bold=True, color=NAVY)
INSIGHT_FONT = Font(name="Calibri", size=11)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def style_header_row(ws, row: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def style_title(ws, row: int, col: int, text: str, merge_to_col: int | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL
    cell.alignment = LEFT
    if merge_to_col and merge_to_col > col:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=merge_to_col)


def apply_alternating_rows(ws, start_row: int, end_row: int, col_count: int) -> None:
    for row in range(start_row, end_row + 1):
        if (row - start_row) % 2 == 1:
            for col in range(1, col_count + 1):
                ws.cell(row=row, column=col).fill = ALT_FILL


def auto_fit_columns(ws, min_width: int = 10, max_width: int = 45) -> None:
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)


def freeze_and_filter(ws, cell: str = "A2") -> None:
    ws.freeze_panes = cell
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def hours_format(cell) -> None:
    cell.number_format = '0.0"h"'


def percent_format(cell) -> None:
    cell.number_format = "0.0%"


def days_format(cell) -> None:
    cell.number_format = '0.0"d"'


def int_format(cell) -> None:
    cell.number_format = "#,##0"
