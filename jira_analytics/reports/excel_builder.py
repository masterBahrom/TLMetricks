"""Build the Phase 5 Excel analytics report with openpyxl."""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from jira_analytics.analytics.models import AnalyticsDocument, Bottleneck
from jira_analytics.config.workflow import WorkflowConfig
from jira_analytics.models.metrics import IssueMetrics, MetricsDocument
from jira_analytics.reports.insights import generate_insights
from jira_analytics.reports.loader import REPORT_VERSION, ReportInputs
from jira_analytics.reports.styles import (
    BODY_FONT,
    BORDER,
    INSIGHT_FONT,
    KPI_BAD,
    KPI_GOOD,
    KPI_LABEL_FONT,
    KPI_VALUE_FONT,
    KPI_WARN,
    LEFT,
    SUBTITLE_FONT,
    apply_alternating_rows,
    auto_fit_columns,
    days_format,
    freeze_and_filter,
    hours_format,
    int_format,
    percent_format,
    style_header_row,
    style_title,
)
from jira_analytics.utils.stats import percentile

SHEET_NAMES = [
    "Executive Summary",
    "Lead Time",
    "Buffer Analysis",
    "Blocked Analysis",
    "Throughput",
    "Flow Analysis",
    "Bug Analysis",
    "Aging",
    "Status Analysis",
    "Reopened Issues",
    "Raw Metrics",
    "Configuration",
    "Management Insights",
]

TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)


class ExcelReportBuilder:
    """Presentation-only Excel workbook from cached metrics and analytics."""

    def __init__(self, inputs: ReportInputs) -> None:
        self.metrics = inputs.metrics
        self.analytics = inputs.analytics
        self.workflow = inputs.workflow
        self.generated_at = inputs.analytics.generated_at

    def build(self, output_path: Path) -> Path:
        wb = Workbook()
        default = wb.active
        wb.remove(default)

        self._build_executive_summary(wb.create_sheet("Executive Summary"))
        lead_ws = wb.create_sheet("Lead Time")
        self._build_lead_time(lead_ws)
        self._build_buffer_analysis(wb.create_sheet("Buffer Analysis"))
        self._build_blocked_analysis(wb.create_sheet("Blocked Analysis"))
        throughput_ws = wb.create_sheet("Throughput")
        self._build_throughput(throughput_ws)
        flow_ws = wb.create_sheet("Flow Analysis")
        self._build_flow_analysis(flow_ws)
        self._build_bug_analysis(wb.create_sheet("Bug Analysis"))
        aging_ws = wb.create_sheet("Aging")
        self._build_aging(aging_ws)
        status_ws = wb.create_sheet("Status Analysis")
        self._build_status_analysis(status_ws)
        self._build_reopened(wb.create_sheet("Reopened Issues"))
        self._build_raw_metrics(wb.create_sheet("Raw Metrics"))
        self._build_configuration(wb.create_sheet("Configuration"))
        self._build_management_insights(wb.create_sheet("Management Insights"))

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return output_path

    # ------------------------------------------------------------------ helpers

    def _reporting_period(self) -> str:
        dates: list[date] = []
        for issue in self.metrics.issues:
            if issue.created_date:
                dates.append(issue.created_date.date())
            end = issue.resolution_date or issue.done_date
            if end:
                dates.append(end.date())
        if not dates:
            return "N/A"
        return f"{min(dates).isoformat()} to {max(dates).isoformat()}"

    def _top_bottleneck(self) -> Bottleneck | None:
        bn = self.analytics.bottlenecks
        candidates = [
            bn.largest_qa,
            bn.largest_review,
            bn.largest_waiting,
            bn.largest_queue,
            bn.largest_active,
        ]
        present = [b for b in candidates if b is not None]
        return max(present, key=lambda b: b.average_seconds) if present else None

    def _bottleneck_label(self, bottleneck: Bottleneck | None) -> str:
        if bottleneck is None:
            return "N/A"
        return f"{bottleneck.status} ({bottleneck.average_hours:.1f}h)"

    def _add_table(self, ws, ref: str, name: str) -> None:
        display_name = "".join(ch for ch in name if ch.isalnum())[:20] or "DataTable"
        table = Table(displayName=display_name, ref=ref)
        table.tableStyleInfo = TABLE_STYLE
        ws.add_table(table)

    def _write_data_table(
        self,
        ws,
        headers: list[str],
        rows: list[list],
        *,
        table_name: str,
        number_cols: dict[int, str] | None = None,
        start_row: int = 1,
    ) -> int:
        """Write header + rows; return last data row index."""
        header_row = start_row
        data_start = start_row + 1
        for col, header in enumerate(headers, start=1):
            ws.cell(row=header_row, column=col, value=header)
        style_header_row(ws, header_row, len(headers))

        for row_idx, row in enumerate(rows, start=data_start):
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = BODY_FONT
                cell.border = BORDER
                if number_cols and col_idx in number_cols:
                    fmt = number_cols[col_idx]
                    if fmt == "hours":
                        hours_format(cell)
                    elif fmt == "percent":
                        percent_format(cell)
                    elif fmt == "days":
                        days_format(cell)
                    elif fmt == "int":
                        int_format(cell)

        last_row = max(header_row, len(rows) + data_start - 1)
        if rows:
            apply_alternating_rows(ws, data_start, last_row, len(headers))
            end_col = get_column_letter(len(headers))
            self._add_table(ws, f"A{header_row}:{end_col}{last_row}", table_name)

        auto_fit_columns(ws)
        if start_row == 1:
            freeze_and_filter(ws)
        return last_row

    def _lead_time_stats(self) -> dict[str, float | None]:
        values = [
            issue.lead_time_hours
            for issue in self.metrics.issues
            if issue.lead_time_hours is not None
        ]
        if not values:
            return {
                "average": None,
                "median": None,
                "p75": None,
                "p90": None,
                "p95": None,
                "maximum": None,
            }
        return {
            "average": sum(values) / len(values),
            "median": percentile(values, 50),
            "p75": percentile(values, 75),
            "p90": percentile(values, 90),
            "p95": percentile(values, 95),
            "maximum": max(values),
        }

    def _status_percentiles(self) -> dict[str, dict[str, float | None]]:
        """Median and P90 per status from per-issue time_in_status (presentation layer)."""
        by_status: dict[str, list[float]] = defaultdict(list)
        for issue in self.metrics.issues:
            for entry in issue.time_in_status:
                by_status[entry.status].append(entry.hours)

        result: dict[str, dict[str, float | None]] = {}
        for status, hours_list in by_status.items():
            result[status] = {
                "median": percentile(hours_list, 50),
                "p90": percentile(hours_list, 90),
                "count": len(hours_list),
            }
        return result

    # -------------------------------------------------------- Executive Summary

    def _build_executive_summary(self, ws) -> None:
        ws.sheet_view.showGridLines = False
        style_title(ws, 1, 1, f"Jira Analytics — {self.analytics.project_key}", merge_to_col=6)

        meta_row = 3
        ws.cell(row=meta_row, column=1, value="Reporting Period").font = KPI_LABEL_FONT
        ws.cell(row=meta_row, column=2, value=self._reporting_period()).font = BODY_FONT
        ws.cell(row=meta_row + 1, column=1, value="Generated At").font = KPI_LABEL_FONT
        ws.cell(row=meta_row + 1, column=2, value=self.generated_at.strftime("%Y-%m-%d %H:%M UTC")).font = BODY_FONT

        summary = self.analytics.summary
        project = self.metrics.project
        aging = self.analytics.aging
        top_bn = self._top_bottleneck()

        kpis: list[tuple[str, object, str]] = [
            ("Project", self.analytics.project_key, "text"),
            ("Total Issues", summary.total_issues, "int"),
            ("Completed Issues", summary.completed_issues, "int"),
            ("Open Issues", summary.open_issues, "int"),
            (
                "Lead Time (median)",
                project.median_lead_time_seconds / 3600 if project.median_lead_time_seconds else None,
                "hours_lower",
            ),
            (
                "Buffer Time (median)",
                project.median_buffer_time_seconds / 3600 if project.median_buffer_time_seconds else None,
                "hours_lower",
            ),
            (
                "Cycle Time (median)",
                project.median_cycle_time_seconds / 3600 if project.median_cycle_time_seconds else None,
                "hours_lower",
            ),
            ("Flow Efficiency", (summary.average_flow_efficiency_percent or 0) / 100, "pct_higher"),
            ("Throughput", summary.total_throughput, "int"),
            (
                "Bug Rate",
                (self.metrics.bug_analytics.bug_rate_percent or project.bug_rate_percent or 0) / 100,
                "pct_lower",
            ),
            ("Reopen %", summary.reopen_percent / 100, "pct_lower"),
            (
                "Average Aging",
                aging.average_seconds / 3600 if aging.average_seconds else None,
                "hours_lower",
            ),
            (
                "P90 Aging",
                aging.p90_seconds / 3600 if aging.p90_seconds else None,
                "hours_lower",
            ),
            ("Top Bottleneck", self._bottleneck_label(top_bn), "text"),
            ("Largest Queue", self._bottleneck_label(self.analytics.bottlenecks.largest_queue), "text"),
            (
                "Largest Waiting Stage",
                self._bottleneck_label(self.analytics.bottlenecks.largest_waiting),
                "text",
            ),
            (
                "Largest Review Stage",
                self._bottleneck_label(self.analytics.bottlenecks.largest_review),
                "text",
            ),
            ("Largest QA Stage", self._bottleneck_label(self.analytics.bottlenecks.largest_qa), "text"),
        ]

        start_row = 7
        for idx, (label, value, kind) in enumerate(kpis):
            row = start_row + idx
            label_cell = ws.cell(row=row, column=2, value=label)
            label_cell.font = KPI_LABEL_FONT
            label_cell.alignment = LEFT

            value_cell = ws.cell(row=row, column=3, value=value if value is not None else "N/A")
            value_cell.font = KPI_VALUE_FONT
            value_cell.alignment = LEFT
            value_cell.border = BORDER

            if kind == "hours_lower" and isinstance(value, (int, float)):
                hours_format(value_cell)
                self._apply_kpi_hours_fill(value_cell, value)
            elif kind == "pct_higher" and isinstance(value, (int, float)):
                percent_format(value_cell)
                self._apply_kpi_percent_higher(value_cell, value)
            elif kind == "pct_lower" and isinstance(value, (int, float)):
                percent_format(value_cell)
                self._apply_kpi_percent_lower(value_cell, value)
            elif kind == "int":
                int_format(value_cell)

        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 36
        ws.column_dimensions["D"].width = 2

    def _apply_kpi_hours_fill(self, cell, hours: float) -> None:
        if hours <= 48:
            cell.fill = KPI_GOOD
        elif hours <= 120:
            cell.fill = KPI_WARN
        else:
            cell.fill = KPI_BAD

    def _apply_kpi_percent_higher(self, cell, pct: float) -> None:
        if pct >= 0.6:
            cell.fill = KPI_GOOD
        elif pct >= 0.4:
            cell.fill = KPI_WARN
        else:
            cell.fill = KPI_BAD

    def _apply_kpi_percent_lower(self, cell, pct: float) -> None:
        if pct <= 0.05:
            cell.fill = KPI_GOOD
        elif pct <= 0.15:
            cell.fill = KPI_WARN
        else:
            cell.fill = KPI_BAD

    # -------------------------------------------------------------- Lead Time

    def _build_lead_time(self, ws) -> None:
        headers = [
            "Issue Key",
            "Summary",
            "Lead Time (h)",
            "Cycle Time (h)",
            "Net Cycle Time (h)",
            "Buffer Time (h)",
            "Blocked Time (h)",
            "Time to First Progress (h)",
            "Waiting Time (h)",
            "Active Time (h)",
            "Reopened",
            "Terminal Status",
        ]
        sorted_issues = sorted(
            self.metrics.issues,
            key=lambda i: i.lead_time_hours or 0,
            reverse=True,
        )
        rows = []
        for issue in sorted_issues:
            rows.append([
                issue.issue_key,
                issue.summary or "—",
                issue.lead_time_hours,
                issue.cycle_time_hours,
                issue.net_cycle_time_hours,
                issue.buffer_time_hours,
                issue.blocked_time_hours,
                issue.time_to_first_progress_hours,
                issue.waiting_time_hours,
                issue.total_active_time_hours,
                "Yes" if issue.is_reopened else "No",
                issue.current_status or "—",
            ])

        last_row = self._write_data_table(
            ws,
            headers,
            rows,
            table_name="LeadTime",
            number_cols={
                3: "hours",
                4: "hours",
                5: "hours",
                6: "hours",
                7: "hours",
                8: "hours",
                9: "hours",
                10: "hours",
            },
        )

        stats = self._lead_time_stats()
        stats_row = last_row + 3
        ws.cell(row=stats_row, column=1, value="Lead Time Statistics").font = Font(bold=True, size=12)
        stat_labels = ["Average", "Median", "P75", "P90", "P95", "Maximum"]
        stat_keys = ["average", "median", "p75", "p90", "p95", "maximum"]
        for col, (label, key) in enumerate(zip(stat_labels, stat_keys), start=1):
            ws.cell(row=stats_row + 1, column=col, value=label).font = KPI_LABEL_FONT
            val = stats[key]
            cell = ws.cell(row=stats_row + 2, column=col, value=val)
            hours_format(cell)

        if rows:
            chart = BarChart()
            chart.type = "col"
            chart.title = "Lead Time Distribution"
            chart.y_axis.title = "Hours"
            chart.x_axis.title = "Issue"
            data = Reference(ws, min_col=3, min_row=1, max_row=last_row)
            cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.width = 18
            chart.height = 10
            ws.add_chart(chart, f"K{stats_row}")

    # -------------------------------------------------------- Buffer Analysis

    def _build_buffer_analysis(self, ws) -> None:
        project = self.metrics.project
        ws.sheet_view.showGridLines = False
        style_title(ws, 1, 1, "Buffer Analysis", merge_to_col=6)

        stats_row = 3
        stat_labels = [
            ("Average Buffer (h)", project.average_buffer_time_seconds),
            ("Median Buffer (h)", project.median_buffer_time_seconds),
            ("P75 Buffer (h)", project.p75_buffer_time_seconds),
            ("P90 Buffer (h)", project.p90_buffer_time_seconds),
            ("P95 Buffer (h)", project.p95_buffer_time_seconds),
        ]
        for idx, (label, seconds) in enumerate(stat_labels):
            row = stats_row + idx
            ws.cell(row=row, column=1, value=label).font = KPI_LABEL_FONT
            value = seconds / 3600 if seconds is not None else None
            cell = ws.cell(row=row, column=2, value=value if value is not None else "N/A")
            if value is not None:
                hours_format(cell)

        headers = ["Issue Key", "Summary", "Buffer (h)", "Lead (h)", "Active (h)", "Flow Efficiency (%)"]
        done_issues = sorted(
            [issue for issue in self.metrics.issues if issue.is_done and issue.buffer_time_seconds > 0],
            key=lambda i: i.buffer_time_seconds,
            reverse=True,
        )[:20]
        rows = [
            [
                issue.issue_key,
                issue.summary or "—",
                issue.buffer_time_hours,
                issue.lead_time_hours,
                issue.total_active_time_hours,
                issue.flow_efficiency_percent,
            ]
            for issue in done_issues
        ]
        table_start = stats_row + len(stat_labels) + 2
        ws.cell(row=table_start - 1, column=1, value="Top 20 Issues by Buffer Time").font = Font(bold=True, size=12)
        self._write_data_table(
            ws,
            headers,
            rows,
            table_name="BufferTop20",
            start_row=table_start,
            number_cols={3: "hours", 4: "hours", 5: "hours", 6: "percent"},
        )

    # ------------------------------------------------------- Blocked Analysis

    def _build_blocked_analysis(self, ws) -> None:
        project = self.metrics.project
        ws.sheet_view.showGridLines = False
        style_title(ws, 1, 1, "Blocked Analysis", merge_to_col=7)

        stat_labels = [
            ("Average Blocked (h)", project.average_blocked_time_seconds),
            ("Median Blocked (h)", project.median_blocked_time_seconds),
            ("P75 Blocked (h)", project.p75_blocked_time_seconds),
            ("P90 Blocked (h)", project.p90_blocked_time_seconds),
            ("P95 Blocked (h)", project.p95_blocked_time_seconds),
            ("Blocked Issue %", project.blocked_issue_percent),
        ]
        for idx, (label, value) in enumerate(stat_labels, start=3):
            ws.cell(row=idx, column=1, value=label).font = KPI_LABEL_FONT
            if value is not None and "Blocked (h)" in label:
                cell_value = value / 3600
            elif value is not None and label.endswith("%"):
                cell_value = value / 100
            else:
                cell_value = value
            cell = ws.cell(row=idx, column=2, value=cell_value if cell_value is not None else "N/A")
            if value is not None:
                percent_format(cell) if label.endswith("%") else hours_format(cell)

        headers = [
            "Issue Key",
            "Summary",
            "Blocked (h)",
            "Terminal Blocked (h)",
            "Lead (h)",
            "Net Cycle (h)",
            "Net Flow Efficiency (%)",
            "Status",
            "Assignee",
        ]
        blocked_issues = sorted(
            [issue for issue in self.metrics.issues if issue.blocked_time_seconds > 0],
            key=lambda i: i.blocked_time_seconds,
            reverse=True,
        )[:20]
        rows = [
            [
                issue.issue_key,
                issue.summary or "—",
                issue.blocked_time_hours,
                issue.terminal_blocked_time_hours,
                issue.lead_time_hours,
                issue.net_cycle_time_hours,
                issue.net_flow_efficiency_percent / 100 if issue.net_flow_efficiency_percent is not None else None,
                issue.current_status or "—",
                issue.assignee or "—",
            ]
            for issue in blocked_issues
        ]
        table_start = 11
        ws.cell(row=table_start - 1, column=1, value="Top 20 Issues by Blocked Time").font = Font(bold=True, size=12)
        self._write_data_table(
            ws,
            headers,
            rows,
            table_name="BlockedTop20",
            start_row=table_start,
            number_cols={3: "hours", 4: "hours", 5: "hours", 6: "hours", 7: "percent"},
        )

    # ----------------------------------------------------------- Bug Analysis

    def _build_bug_analysis(self, ws) -> None:
        bugs = self.metrics.bug_analytics
        ws.sheet_view.showGridLines = False
        style_title(ws, 1, 1, "Bug Analysis", merge_to_col=6)

        summary = [
            ("Completed Bugs", bugs.completed_bugs),
            ("Completed HOT-FIX", bugs.completed_hotfix),
            ("Total Bugs", bugs.total_bugs),
            ("Completed Issues", bugs.completed_issues),
            ("Bug Rate (%)", bugs.bug_rate_percent),
        ]
        row = 3
        for label, value in summary:
            ws.cell(row=row, column=1, value=label).font = KPI_LABEL_FONT
            ws.cell(row=row, column=2, value=value if value is not None else "N/A").font = BODY_FONT
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="By Sprint").font = Font(bold=True, size=12)
        row += 1
        sprint_headers = ["Sprint", "Completed Bugs", "Completed Issues", "Bug Rate (%)"]
        sprint_rows = [
            [entry.sprint, entry.completed_bugs, entry.completed_issues, entry.bug_rate_percent]
            for entry in bugs.by_sprint
        ]
        row = self._write_data_table(
            ws, sprint_headers, sprint_rows, table_name="BugBySprint", start_row=row, number_cols={4: "percent"}
        )

        row += 2
        ws.cell(row=row, column=1, value="By Issue Type").font = Font(bold=True, size=12)
        row += 1
        type_headers = ["Issue Type", "Completed", "Bug Count", "Bug %"]
        type_rows = []
        for entry in bugs.by_issue_type:
            pct = (entry.bug_count / entry.completed_count * 100) if entry.completed_count else 0
            type_rows.append([entry.issue_type, entry.completed_count, entry.bug_count, pct])
        self._write_data_table(
            ws, type_headers, type_rows, table_name="BugByType", start_row=row, number_cols={4: "percent"}
        )

    # ------------------------------------------------------------- Throughput

    def _build_throughput(self, ws) -> None:
        headers = ["Granularity", "Period", "Count"]
        rows: list[list] = []
        for label, buckets in [
            ("Daily", self.analytics.throughput.daily),
            ("Weekly", self.analytics.throughput.weekly),
            ("Monthly", self.analytics.throughput.monthly),
        ]:
            for bucket in buckets:
                rows.append([label, bucket.period_label, bucket.count])

        total = self.analytics.summary.total_throughput
        weekly = self.analytics.throughput.weekly
        monthly = self.analytics.throughput.monthly
        avg_week = sum(b.count for b in weekly) / len(weekly) if weekly else 0
        avg_month = sum(b.count for b in monthly) / len(monthly) if monthly else 0

        last_row = self._write_data_table(ws, headers, rows, table_name="Throughput", number_cols={3: "int"})

        summary_row = last_row + 3
        ws.cell(row=summary_row, column=1, value="Totals").font = Font(bold=True)
        ws.cell(row=summary_row + 1, column=1, value="Total Throughput")
        ws.cell(row=summary_row + 1, column=2, value=total)
        ws.cell(row=summary_row + 2, column=1, value="Average per Week")
        c = ws.cell(row=summary_row + 2, column=2, value=round(avg_week, 2))
        int_format(c)
        ws.cell(row=summary_row + 3, column=1, value="Average per Month")
        c = ws.cell(row=summary_row + 3, column=2, value=round(avg_month, 2))
        int_format(c)

        daily_count = len(self.analytics.throughput.daily)
        weekly_start = 2 + daily_count
        weekly_end = weekly_start + len(self.analytics.throughput.weekly) - 1
        if weekly_end >= weekly_start:
            w_chart = BarChart()
            w_chart.title = "Weekly Throughput"
            w_chart.y_axis.title = "Issues"
            w_data = Reference(ws, min_col=3, min_row=weekly_start, max_row=weekly_end)
            w_cats = Reference(ws, min_col=2, min_row=weekly_start, max_row=weekly_end)
            w_chart.add_data(w_data, titles_from_data=False)
            w_chart.set_categories(w_cats)
            ws.add_chart(w_chart, "E2")

        monthly_start = 2 + daily_count + len(self.analytics.throughput.weekly)
        monthly_end = monthly_start + len(self.analytics.throughput.monthly) - 1
        if monthly_end >= monthly_start:
            m_chart = BarChart()
            m_chart.title = "Monthly Throughput"
            m_chart.y_axis.title = "Issues"
            m_data = Reference(ws, min_col=3, min_row=monthly_start, max_row=monthly_end)
            m_cats = Reference(ws, min_col=2, min_row=monthly_start, max_row=monthly_end)
            m_chart.add_data(m_data, titles_from_data=False)
            m_chart.set_categories(m_cats)
            ws.add_chart(m_chart, "E18")

    # ---------------------------------------------------------- Flow Analysis

    def _build_flow_analysis(self, ws) -> None:
        headers = ["Metric", "Value"]
        flow = self.analytics.flow
        queues = {q.category: q for q in self.analytics.queues.categories}

        def pct(cat: str) -> float | None:
            entry = queues.get(cat)
            return entry.percent_of_total / 100 if entry else None

        rows = [
            ["Flow Efficiency (avg %)", (flow.average_efficiency_percent or 0) / 100],
            ["Flow Efficiency (median %)", (flow.median_efficiency_percent or 0) / 100],
            ["Queue %", pct("queue")],
            ["Waiting %", pct("waiting")],
            ["QA %", pct("qa")],
            ["Review %", pct("review")],
            ["Active %", pct("active")],
            ["Done %", pct("done")],
        ]
        last_row = self._write_data_table(
            ws,
            headers,
            rows,
            table_name="FlowMetrics",
            number_cols={2: "percent"},
        )

        bn_row = last_row + 3
        ws.cell(row=bn_row, column=1, value="Largest Bottlenecks").font = Font(bold=True, size=12)
        bn_headers = ["Category", "Status", "Average Hours"]
        for col, h in enumerate(bn_headers, start=1):
            ws.cell(row=bn_row + 1, column=col, value=h)
        style_header_row(ws, bn_row + 1, len(bn_headers))

        bn = self.analytics.bottlenecks
        bn_rows = [
            ("Queue", bn.largest_queue),
            ("Waiting", bn.largest_waiting),
            ("Review", bn.largest_review),
            ("QA", bn.largest_qa),
            ("Active", bn.largest_active),
        ]
        for idx, (cat, bottleneck) in enumerate(bn_rows, start=bn_row + 2):
            ws.cell(row=idx, column=1, value=cat)
            ws.cell(row=idx, column=2, value=bottleneck.status if bottleneck else "N/A")
            cell = ws.cell(row=idx, column=3, value=bottleneck.average_hours if bottleneck else None)
            hours_format(cell)

        pie = PieChart()
        pie.title = "Flow Efficiency Mix"
        pie_data = Reference(ws, min_col=2, min_row=3, max_row=10)
        pie_labels = Reference(ws, min_col=1, min_row=3, max_row=10)
        pie.add_data(pie_data, titles_from_data=False)
        pie.set_categories(pie_labels)
        ws.add_chart(pie, "E2")

    # ------------------------------------------------------------------ Aging

    def _build_aging(self, ws) -> None:
        headers = [
            "Issue",
            "Status",
            "Assignee",
            "Created",
            "Current Status Since",
            "Aging Days",
        ]
        metrics_by_key = {issue.issue_key: issue for issue in self.metrics.issues}
        aging_entries = sorted(self.analytics.aging.top_oldest, key=lambda e: e.aging_days, reverse=True)

        rows = []
        for entry in aging_entries:
            issue = metrics_by_key.get(entry.issue_key)
            created = issue.created_date.strftime("%Y-%m-%d") if issue and issue.created_date else "N/A"
            rows.append([
                entry.issue_key,
                entry.current_status,
                issue.assignee or "—" if issue else "—",
                created,
                "N/A",
                entry.aging_days,
            ])

        if not rows:
            rows.append(["No open issues", "—", "—", "—", "—", 0])

        last_row = self._write_data_table(
            ws,
            headers,
            rows,
            table_name="Aging",
            number_cols={6: "days"},
        )

        if self.analytics.aging.open_issue_count > 0:
            hist_start = last_row + 3
            ws.cell(row=hist_start, column=1, value="Aging Bucket").font = KPI_LABEL_FONT
            ws.cell(row=hist_start, column=2, value="Count").font = KPI_LABEL_FONT
            buckets = [("0-7d", 0), ("8-14d", 0), ("15-30d", 0), ("31-60d", 0), ("60+d", 0)]
            for entry in aging_entries:
                days = entry.aging_days
                if days <= 7:
                    buckets[0] = (buckets[0][0], buckets[0][1] + 1)
                elif days <= 14:
                    buckets[1] = (buckets[1][0], buckets[1][1] + 1)
                elif days <= 30:
                    buckets[2] = (buckets[2][0], buckets[2][1] + 1)
                elif days <= 60:
                    buckets[3] = (buckets[3][0], buckets[3][1] + 1)
                else:
                    buckets[4] = (buckets[4][0], buckets[4][1] + 1)

            for idx, (label, count) in enumerate(buckets, start=hist_start + 1):
                ws.cell(row=idx, column=1, value=label)
                ws.cell(row=idx, column=2, value=count)

            chart = BarChart()
            chart.title = "Aging Histogram"
            chart.y_axis.title = "Issues"
            data = Reference(ws, min_col=2, min_row=hist_start, max_row=hist_start + len(buckets))
            cats = Reference(ws, min_col=1, min_row=hist_start + 1, max_row=hist_start + len(buckets))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, "H2")

    # --------------------------------------------------------- Status Analysis

    def _build_status_analysis(self, ws) -> None:
        headers = [
            "Status",
            "Average Time (h)",
            "Count",
            "Median (h)",
            "P90 (h)",
            "Status Category",
        ]
        percentiles = self._status_percentiles()
        dist = self.analytics.distributions.status.by_status

        rows = []
        for entry in sorted(dist, key=lambda e: e.average_hours, reverse=True):
            stats = percentiles.get(entry.status, {})
            rows.append([
                entry.status,
                entry.average_hours,
                entry.issue_count,
                stats.get("median"),
                stats.get("p90"),
                entry.workflow_type,
            ])

        if not rows:
            for avg in self.metrics.project.average_time_by_status:
                stats = percentiles.get(avg.status, {})
                rows.append([
                    avg.status,
                    avg.average_hours,
                    avg.issue_count,
                    stats.get("median"),
                    stats.get("p90"),
                    self.workflow.queue_category(avg.status),
                ])

        last_row = self._write_data_table(
            ws,
            headers,
            rows,
            table_name="StatusAnalysis",
            number_cols={2: "hours", 4: "hours", 5: "hours"},
        )

        if rows:
            chart = BarChart()
            chart.title = "Status Distribution"
            chart.y_axis.title = "Average Hours"
            data = Reference(ws, min_col=2, min_row=1, max_row=last_row)
            cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.width = 16
            chart.height = 10
            ws.add_chart(chart, "H2")

    # -------------------------------------------------------- Reopened Issues

    def _build_reopened(self, ws) -> None:
        headers = ["Issue", "Summary", "Reopen Count", "Lead Time (h)", "Current Status"]
        reopened = [
            issue
            for issue in self.metrics.issues
            if issue.is_reopened or issue.reopen_count > 0
        ]
        rows = [
            [
                issue.issue_key,
                issue.summary or "—",
                issue.reopen_count,
                issue.lead_time_hours,
                issue.current_status or "—",
            ]
            for issue in sorted(reopened, key=lambda i: i.reopen_count, reverse=True)
        ]
        if not rows:
            rows.append(["No reopened issues", "—", 0, None, "—"])

        self._write_data_table(
            ws,
            headers,
            rows,
            table_name="Reopened",
            number_cols={3: "int", 4: "hours"},
        )

    # ------------------------------------------------------------- Raw Metrics

    def _build_raw_metrics(self, ws) -> None:
        if not self.metrics.issues:
            self._write_data_table(ws, ["issue_key"], [["No data"]], table_name="RawMetrics")
            return

        sample = self.metrics.issues[0].model_dump(mode="json")
        flat_keys = list(sample.keys())
        if sample.get("time_in_status"):
            flat_keys = [k for k in flat_keys if k != "time_in_status"]
            flat_keys.append("time_in_status_json")
        for list_key in ("labels", "components"):
            if list_key in flat_keys:
                flat_keys.remove(list_key)
                flat_keys.append(f"{list_key}_json")

        headers = flat_keys
        rows = []
        for issue in self.metrics.issues:
            data = issue.model_dump(mode="json")
            row = []
            for key in flat_keys:
                if key == "time_in_status_json":
                    row.append(json.dumps(data.get("time_in_status", [])))
                elif key == "labels_json":
                    row.append(json.dumps(data.get("labels", [])))
                elif key == "components_json":
                    row.append(json.dumps(data.get("components", [])))
                else:
                    value = data.get(key)
                    if isinstance(value, (list, dict)):
                        row.append(json.dumps(value))
                    else:
                        row.append(value)
            rows.append(row)

        self._write_data_table(ws, headers, rows, table_name="RawMetrics")

    # ----------------------------------------------------------- Configuration

    def _build_configuration(self, ws) -> None:
        ws.sheet_view.showGridLines = False
        style_title(ws, 1, 1, "Workflow Configuration", merge_to_col=4)

        sections = [
            ("Terminal Statuses", self.workflow.terminal_statuses),
            ("Queue / Buffer Statuses", self.workflow.queue_statuses),
            ("Active Statuses", self.workflow.active_statuses),
            ("Waiting Statuses", self.workflow.waiting_statuses),
            ("QA Statuses", self.workflow.qa_statuses),
            ("Review Statuses", self.workflow.review_statuses),
            ("Deployment Queue Statuses", self.workflow.deployment_queue_statuses),
            ("Cancelled Statuses", self.workflow.cancelled_statuses),
            ("Bug Issue Types", self.workflow.bug_issue_types),
        ]

        row = 3
        for title, statuses in sections:
            ws.cell(row=row, column=1, value=title).font = Font(bold=True)
            row += 1
            for status in statuses:
                ws.cell(row=row, column=2, value=status).font = BODY_FONT
                row += 1
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Metadata").font = Font(bold=True, size=12)
        meta = [
            ("Project", self.analytics.project_key),
            ("Generated Timestamp", self.generated_at.isoformat()),
            ("Report Version", REPORT_VERSION),
            ("Metrics Generated", self.metrics.generated_at.isoformat()),
            ("Analytics Generated", self.analytics.generated_at.isoformat()),
        ]
        for idx, (label, value) in enumerate(meta, start=row + 1):
            ws.cell(row=idx, column=1, value=label).font = KPI_LABEL_FONT
            ws.cell(row=idx, column=2, value=value).font = BODY_FONT

        auto_fit_columns(ws)

    # ---------------------------------------------------- Management Insights

    def _build_management_insights(self, ws) -> None:
        ws.sheet_view.showGridLines = False
        style_title(ws, 1, 1, "Management Insights", merge_to_col=6)
        ws.cell(row=3, column=1, value="Automatically generated observations").font = SUBTITLE_FONT

        insights = generate_insights(self.metrics, self.analytics, self.workflow)
        for idx, text in enumerate(insights, start=5):
            cell = ws.cell(row=idx, column=1, value=f"• {text}")
            cell.font = INSIGHT_FONT
            cell.alignment = Alignment(wrap_text=True)
            ws.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=6)

        ws.column_dimensions["A"].width = 90
